"""Turn a file (TXT, MD, PDF, DOCX) into a Document with text and metadata."""
import re
from datetime import datetime, timezone
from pathlib import Path

import docx
from pypdf import PdfReader

from app.models import Document

SUPPORTED_FORMATS = {".txt", ".md", ".pdf", ".docx"}


class UnsupportedFormatError(ValueError):
    pass


def make_doc_id(name: str) -> str:
    """Stable id from the file name, so uploading the same file again means an update."""
    stem = Path(name).stem.lower()
    return "file::" + re.sub(r"[^a-z0-9]+", "-", stem).strip("-")


def load_file(path: Path, original_name: str | None = None) -> Document:
    name = original_name or path.name
    suffix = Path(name).suffix.lower()
    if suffix not in SUPPORTED_FORMATS:
        raise UnsupportedFormatError(f"{suffix or 'no extension'} is not supported, use {sorted(SUPPORTED_FORMATS)}")
    readers = {".txt": _read_text, ".md": _read_markdown, ".pdf": _read_pdf, ".docx": _read_docx}
    text, title, extra = readers[suffix](path)
    text = _normalize(text)
    if not text:
        raise ValueError(f"{name}: no text could be extracted")
    metadata = {"file_name": name, "loaded_at": datetime.now(timezone.utc).isoformat(), **extra}
    return Document(
        doc_id=make_doc_id(name),
        title=title or Path(name).stem,
        text=text,
        source=name,
        fmt=suffix.lstrip("."),
        metadata=metadata,
    )


def _normalize(text: str) -> str:
    """Unify line breaks, drop trailing spaces, keep blank lines between paragraphs."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [line.rstrip() for line in text.split("\n")]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def _read_text(path: Path) -> tuple[str, str | None, dict]:
    return path.read_text(encoding="utf-8", errors="replace"), None, {}


def _read_markdown(path: Path) -> tuple[str, str | None, dict]:
    text = path.read_text(encoding="utf-8", errors="replace")
    heading = re.search(r"^#\s+(.+)$", text, flags=re.MULTILINE)
    return text, heading.group(1).strip() if heading else None, {}


def _read_pdf(path: Path) -> tuple[str, str | None, dict]:
    reader = PdfReader(path)
    pages = [page.extract_text() or "" for page in reader.pages]
    info = reader.metadata
    title = info.title if info and info.title else None
    return "\n\n".join(pages), title, {"pages": len(reader.pages)}


def _read_docx(path: Path) -> tuple[str, str | None, dict]:
    document = docx.Document(str(path))
    paragraphs = [p.text for p in document.paragraphs]
    props = document.core_properties
    extra = {"author": props.author} if props.author else {}
    return "\n\n".join(paragraphs), props.title or None, extra
